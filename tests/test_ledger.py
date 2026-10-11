import io
from decimal import Decimal
import pytest
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas
from conftest import row,upload,csv_bytes
import ledger


def transactions(setup):
    return setup[1].get('/api/data').json['transactions']


def test_duplicate_overlap_and_identical_occurrences(setup):
    first=[row(),row(),row(transaction_date='2026-02-01',amount='3.00')]
    assert upload(setup,first)['added']==3
    assert upload(setup,first)['duplicate_file']
    result=upload(setup,[row(),row(),row(transaction_date='2026-02-02',amount='5.00')])
    assert result=={'added':1,'duplicates':2,'updated':0}
    assert len(transactions(setup))==4
    assert ledger.summarize(transactions(setup))['totals']['expense']==2820


def test_explicit_bank_ids_distinguish_same_amount_description(setup):
    upload(setup,[row(external_id='bank-a'),row(external_id='bank-b')])
    assert len(transactions(setup))==2


def test_repeated_bank_id_within_one_file(setup):
    result=upload(setup,[row(external_id='same'),row(external_id='same')])
    assert result=={'added':1,'duplicates':1,'updated':0}
    assert len(transactions(setup)[0]['sources'])==2


def test_month_year_boundaries_posting_date_and_exact_money(setup):
    upload(setup,[row(transaction_date='2025-12-31',posting_date='2026-01-02',amount='0.10'),
                  row(amount='0.20'),row(transaction_date='2026-02-01',amount='0.30')])
    ts=transactions(setup)
    assert ledger.summarize(ts)['totals']['expense']==60
    assert ledger.summarize([t for t in ts if t['transaction_date'].startswith('2026-01')])['totals']['expense']==20
    assert ledger.summarize([t for t in ts if t['transaction_date'].startswith('2025')])['totals']['expense']==10


def test_financial_types_and_pending(setup):
    rows=[row(description=kind,type=kind,amount='100.00') for kind in ledger.KINDS]
    rows.append(row(status='pending',amount='9.90'))
    upload(setup,rows)
    summary=ledger.summarize(transactions(setup))
    assert summary['totals']['income']==10000
    assert summary['totals']['expense']==10000
    assert summary['totals']['balance']==0
    assert summary['totals']['pending']==990
    for kind in ['payment','transfer','refund','reimbursement','withdrawal']:
        assert summary['totals'][kind]==10000


def test_pending_to_statement_preserves_id_manual_notes_rules(setup):
    _,client,headers,_=setup
    upload(setup,[row(status='pending',external_id='bank-1')])
    tid=transactions(setup)[0]['id']
    assert client.patch('/api/transactions/'+tid,headers=headers,json={'category':'学习','subcategory':'书籍','notes':'个人用途', 'future_rule':True}).status_code==200
    result=upload(setup,[row(external_id='bank-1',posting_date='2026-02-03')])
    assert result['updated']==1
    current=transactions(setup)[0]
    assert (current['id'],current['category'],current['notes'],current['status'])==(tid,'学习','个人用途','posted')
    assert len(current['sources'])==2
    upload(setup,[row(transaction_date='2026-02-03',external_id='bank-2')])
    assert all(t['category']=='学习' for t in transactions(setup))
    assert ledger.summarize(transactions(setup))['categories']=={'学习':2020}


def test_explicit_changed_pending_match_alias_repeat(setup):
    upload(setup,[row(status='pending',transaction_date='2026-01-30',description='Temporary')])
    tid=transactions(setup)[0]['id']
    formal=row(transaction_date='2026-02-01',description='Settled')
    assert upload(setup,[formal],decisions={'0':{'match_id':tid,'confirm':True}})['updated']==1
    upload(setup,[formal,row(transaction_date='2026-02-02',amount='5.00')])
    assert len(transactions(setup))==2
    assert next(t for t in transactions(setup) if t['id']==tid)['transaction_date']=='2026-02-01'


def test_changed_external_id_requires_confirmation(setup):
    upload(setup,[row(external_id='bank-a')])
    _,client,headers,account=setup
    p=client.post('/api/import/preview',headers=headers,data={'account_id':account,'file':(io.BytesIO(csv_bytes([row(external_id='bank-a',amount='11.00')])),'x.csv')}).json
    assert p['rows'][0]['action']=='review'
    assert client.post('/api/import/commit',headers=headers,json={'token':p['token']}).status_code==400
    assert transactions(setup)[0]['amount']==1010
    assert client.post('/api/import/commit',headers=headers,json={'token':p['token'],'decisions':{'0':{'confirm':True}}}).status_code==200
    assert transactions(setup)[0]['amount']==1100


def test_account_isolation_and_currency(setup):
    _,client,headers,account=setup
    upload(setup,[row()])
    client.post('/api/accounts',headers=headers,json={'name':'Fictional savings','kind':'savings','currency':'CAD'})
    second=next(a['id'] for a in client.get('/api/data').json['accounts'] if a['id']!=account)
    upload((setup[0],client,headers,second),[row()])
    assert len(transactions(setup))==2
    p=client.post('/api/import/preview',headers=headers,data={'account_id':account,'file':(io.BytesIO(csv_bytes([row(currency='USD')])),'x.csv')})
    assert p.status_code==400


def test_missing_dates_and_invalid_amount_fail_without_state(setup):
    _,client,headers,account=setup
    for change in [{'transaction_date':''},{'amount':'0.001'},{'amount':'NaN'},{'amount':'-3.00'},{'type':'unknown'},{'posting_date':'01/02/26'}]:
        p=client.post('/api/import/preview',headers=headers,data={'account_id':account,'file':(io.BytesIO(csv_bytes([row(**change)])),'x.csv')})
        assert p.status_code==400,p.json
    assert not transactions(setup)


def test_xlsx_and_pdf_local_parser(setup):
    r=row(external_id='xlsx-1')
    wb=Workbook();ws=wb.active;ws.append(ledger.HEADERS);ws.append([r.get(k,'') for k in ledger.HEADERS]);buf=io.BytesIO();wb.save(buf)
    assert upload(setup,[],filename='fiction.xlsx',raw=buf.getvalue())['added']==1
    pdf=io.BytesIO();c=canvas.Canvas(pdf,pagesize=(1800,600));c.setFont('Helvetica',9)
    c.drawString(20,560,'|'.join(ledger.HEADERS));r=row(external_id='pdf-1',amount='13.23');c.drawString(20,530,'|'.join(str(r.get(k,'')) for k in ledger.HEADERS));c.save()
    assert upload(setup,[],filename='fiction.pdf',raw=pdf.getvalue())['added']==1
    source=next(t for t in transactions(setup) if t['external_id']=='pdf-1')['sources'][0]
    assert source['page']==1 and source['line']==2


def test_hierarchy_excel_formulas_and_text_injection(setup):
    upload(setup,[row(category='饮食',subcategory='咖啡',notes='=HYPERLINK("https://invalid.example")'),
                  row(category='饮食',subcategory='超市',description='Fictional groceries',amount='90.00'),
                  row(category='车',subcategory='停车',description='Fictional parking',amount='5.00')])
    summary=ledger.summarize(transactions(setup))
    assert sum(summary['categories'].values())==summary['totals']['expense']==10510
    for category,amount in summary['categories'].items():
        assert sum(summary['subcategories'][category].values())==amount
    output=setup[1].get('/api/export.xlsx');assert output.status_code==200
    ws=load_workbook(io.BytesIO(output.data),data_only=False)['2026']
    assert ws['R2'].value.startswith('=SUMIFS(') and ws['T2'].value=='=R2-S2'
    assert ws['T3'].value=='=IFERROR(S3/S2,0)'
    notes=[ws.cell(i,11) for i in range(2,5)]
    assert next(c for c in notes if c.value).data_type=='s'
    assert sum(Decimal(str(ws.cell(i,6).value)) for i in range(2,5))==Decimal('105.10')


def test_coverage_explicit_not_inferred(setup):
    upload(setup,[row()],coverage=('2026-01-01','2026-01-31'))
    f=setup[1].get('/api/data').json['files'][0]
    assert (f['coverage_start'],f['coverage_end'])==('2026-01-01','2026-01-31')
    upload(setup,[row(transaction_date='2026-02-28')])
    assert sum(f['coverage_start'] is None for f in setup[1].get('/api/data').json['files'])==1


def test_restart_persistence(setup):
    from app import create_app
    upload(setup,[row()])
    app=create_app(setup[0].config['DATA_DIR'],testing=True)
    with ledger.connect(app.config['DB']) as conn:
        assert conn.execute('SELECT COUNT(*) FROM transactions').fetchone()[0]==1


def test_no_login_page_api_file_export_and_revoked_session(setup):
    app,client,headers,_=setup
    upload(setup,[row()])
    file_id=transactions(setup)[0]['sources'][0]['file_id']
    guest=app.test_client()
    assert guest.get('/').status_code==302
    for path in ['/api/data','/api/export.xlsx','/files/'+file_id,'/api/import/preview']:
        response=guest.post(path) if path.endswith('preview') else guest.get(path)
        assert response.status_code==401
        assert b'Fictional coffee' not in response.data
    assert client.get('/files/'+file_id).status_code==200
    assert client.post('/api/accounts',json={}).status_code==403
    old_cookie=client.get_cookie('session').value
    assert client.post('/logout',headers=headers).status_code==302
    guest.set_cookie('session',old_cookie)
    assert guest.get('/api/data').status_code==401


def test_favicon_does_not_invalidate_login_csrf(setup):
    import re
    app=setup[0]
    guest=app.test_client()
    csrf=re.search(r'name="csrf" value="([^"]+)"',guest.get('/login').text)[1]
    assert guest.get('/favicon.ico').status_code==204
    assert guest.post('/login',data={'csrf':csrf,'password':'only-for-fictional-tests'}).status_code==302
